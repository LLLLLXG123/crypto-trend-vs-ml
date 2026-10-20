"""python-pptx 的小工具：统一字体、标题、要点、表格、图片、页脚和讲稿。"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

FONT = "Microsoft YaHei"
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ALERT = RGBColor(0xC0, 0x00, 0x00)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def style_run(run, size=18, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        latin = rpr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(ea)
        else:
            rpr.append(ea)
    ea.set("typeface", FONT)


def text_box(slide, left, top, width, height, paragraphs, size=18, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=6):
    """paragraphs：字符串，或 [(文字, {size, bold, color, level}), ...]。
    文字里用 **加粗** 标记强调。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(paragraphs, str):
        paragraphs = [(paragraphs, {})]
    for i, item in enumerate(paragraphs):
        text, opts = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(opts.get("space_after", spacing))
        p.level = opts.get("level", 0)
        parts = text.split("**")
        for j, part in enumerate(parts):
            if not part:
                continue
            run = p.add_run()
            run.text = part
            style_run(run, size=opts.get("size", size), bold=opts.get("bold", False) or j % 2 == 1,
                      color=opts.get("color", color if j % 2 == 0 else opts.get("em", ACCENT)))
    return box


def title(slide, text, kicker=None):
    if kicker:
        text_box(slide, MARGIN, Inches(0.32), Inches(12), Inches(0.35), kicker, size=12, color=ACCENT)
    text_box(slide, MARGIN, Inches(0.55), Inches(12.1), Inches(0.8), [(text, {"bold": True})], size=28)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.3), Inches(0.9), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def panel(slide, left, top, width, height, heading, body, size=15):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.fill.background()
    text_box(slide, left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.5),
             [(heading, {"bold": True, "color": ACCENT})], size=18)
    text_box(slide, left + Inches(0.2), top + Inches(0.7), width - Inches(0.4), height - Inches(0.8),
             body, size=size, spacing=8)


def callout(slide, left, top, width, height, text, size=17):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    text_box(slide, left + Inches(0.25), top, width - Inches(0.5), height,
             [(text, {"color": WHITE, "bold": True, "em": WHITE})], size=size, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, left, top, width, rows, col_widths=None, size=12, row_height=0.42, header_size=None,
          left_cols=(0,)):
    """left_cols：左对齐的列（默认只有第一列），其余居中。"""
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(row_height * n_rows))
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = int(width * w / total)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_height)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c in left_cols else PP_ALIGN.CENTER
            text = str(value)
            bold = r == 0 or text.startswith("**")
            run = p.add_run()
            run.text = text.replace("**", "")
            style_run(run, size=header_size or size if r == 0 else size, bold=bold,
                      color=WHITE if r == 0 else INK)
    return tbl


def image(slide, path, left, top, width=None, height=None):
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def footer(slide, page, source):
    text_box(slide, MARGIN, Inches(7.02), Inches(10.5), Inches(0.3), source, size=10, color=MUTED)
    text_box(slide, Inches(12.0), Inches(7.02), Inches(0.75), Inches(0.3), str(page), size=10,
             color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
